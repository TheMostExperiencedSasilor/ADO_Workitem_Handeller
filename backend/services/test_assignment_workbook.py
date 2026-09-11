import posixpath
import re
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from services.ado_client import AdoClient


WORKBOOK_HEADERS = [
    "ID",
    "Title",
    "Product Area",
    "Automation Script Name",
    "Round 1 results",
    "Round 2 results",
    "Single run results",
    "Manual run",
    "Comment",
    "Solution",
    "Defects",
]

RESULT_KEYS = {
    "round1Results": "Round 1 results",
    "round2Results": "Round 2 results",
    "singleRunResults": "Single run results",
    "manualRun": "Manual run",
}


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _field_value(fields: dict[str, Any], aliases: set[str]) -> str:
    normalized_aliases = {_normalize(alias) for alias in aliases}
    for reference_name, value in fields.items():
        tail = reference_name.rsplit(".", 1)[-1]
        normalized_tail = _normalize(tail)
        normalized_full = _normalize(reference_name)
        if normalized_tail in normalized_aliases or any(alias in normalized_full for alias in normalized_aliases):
            if value is None:
                return ""
            if isinstance(value, dict):
                return str(value.get("displayName") or value.get("name") or value.get("value") or "")
            return str(value)
    return ""


def _excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class TestAssignmentWorkbookService:
    def __init__(self, client: AdoClient) -> None:
        self.client = client

    def read_define_cases(self, plan_id: int, suite_id: int) -> list[dict[str, Any]]:
        path = f"Plans/{plan_id}/Suites/{suite_id}/TestCase"
        memberships = self.client._read_test_plan_pages(path)
        if not memberships:
            return []

        case_ids: list[int] = []
        order_by_id: dict[int, int | None] = {}
        for membership in memberships:
            work_item = membership.get("workItem") or {}
            if not work_item.get("id"):
                continue
            case_id = int(work_item["id"])
            if case_id not in order_by_id:
                case_ids.append(case_id)
            order = membership.get("order")
            order_by_id[case_id] = order if isinstance(order, int) and not isinstance(order, bool) else None

        work_items: dict[int, dict[str, Any]] = {}
        for start in range(0, len(case_ids), 200):
            for item in self.client.read_work_items(case_ids[start:start + 200]):
                work_items[int(item["id"])] = item

        rows: list[dict[str, Any]] = []
        for case_id in case_ids:
            fields = work_items.get(case_id, {}).get("fields", {})
            rows.append({
                "testCaseId": case_id,
                "title": str(fields.get("System.Title") or "Title unavailable"),
                "productArea": _field_value(fields, {"ProductArea", "Product Area"}),
                "automationScriptName": _field_value(fields, {
                    "AutomationScriptName",
                    "Automation Script Name",
                    "AutomationScript",
                    "AutomatedTestName",
                }),
                "order": order_by_id.get(case_id),
            })

        known_order = sorted((row for row in rows if row["order"] is not None), key=lambda row: row["order"])
        ordered_iter = iter(known_order)
        return [next(ordered_iter) if row["order"] is not None else row for row in rows]

    def _read_execute_points(self, plan_id: int, suite_id: int) -> list[dict[str, Any]]:
        path = f"Plans/{plan_id}/Suites/{suite_id}/TestPoint"
        points = self.client._read_test_plan_pages(path)
        rows: list[dict[str, Any]] = []
        for point in points:
            reference = point.get("testCaseReference") or {}
            if not reference.get("id"):
                continue
            tester = point.get("tester") or {}
            if isinstance(tester, dict):
                display_name = str(tester.get("displayName") or "").strip()
                unique_name = str(tester.get("uniqueName") or "").strip()
            else:
                display_name = str(tester).strip()
                unique_name = display_name
            rows.append({
                "testCaseId": int(reference["id"]),
                "tester": display_name or unique_name,
                "testerAliases": [name for name in (display_name, unique_name) if name],
            })
        return rows

    def build_assignment_rows(self, plan_id: int, suite_id: int, tester_query: str) -> dict[str, Any]:
        query = tester_query.strip()
        if not query:
            raise ValueError("Assigned tester is required.")

        define_cases = self.read_define_cases(plan_id, suite_id)
        execute_points = self._read_execute_points(plan_id, suite_id)
        query_folded = query.casefold()

        exact_points = [
            point for point in execute_points
            if any(alias.casefold() == query_folded for alias in point["testerAliases"])
        ]
        candidates = exact_points or [
            point for point in execute_points
            if any(query_folded in alias.casefold() for alias in point["testerAliases"])
        ]

        matched_names = list(dict.fromkeys(point["tester"] for point in candidates if point["tester"]))
        if not exact_points and len({name.casefold() for name in matched_names}) > 1:
            raise ValueError(
                "Tester search matches multiple people: " + ", ".join(matched_names[:8])
                + ". Enter the exact display name or email."
            )

        assigned_case_ids = {point["testCaseId"] for point in candidates}
        assigned_rows: list[dict[str, Any]] = []
        for case in define_cases:
            if case["testCaseId"] not in assigned_case_ids:
                continue
            assigned_rows.append({
                **case,
                "tester": matched_names[0] if matched_names else query,
                "round1Results": "",
                "round2Results": "",
                "singleRunResults": "",
                "manualRun": "",
                "comment": "",
                "solution": "",
                "defects": "",
            })

        return {
            "defineCount": len(define_cases),
            "assignedCount": len(assigned_rows),
            "matchedTester": matched_names[0] if len(matched_names) == 1 else (query if assigned_rows else ""),
            "rows": assigned_rows,
        }

    @staticmethod
    def make_workbook(rows: list[dict[str, Any]]) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Main"
        sheet.append(WORKBOOK_HEADERS)

        for row in rows:
            values: list[Any] = [
                row.get("testCaseId"),
                row.get("title", ""),
                row.get("productArea", ""),
                row.get("automationScriptName", ""),
                row.get("round1Results", ""),
                row.get("round2Results", ""),
                row.get("singleRunResults", ""),
                row.get("manualRun", ""),
                row.get("comment", ""),
                row.get("solution", ""),
                row.get("defects", ""),
            ]
            sheet.append([_excel_safe(value) for value in values])

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32

        widths = [12, 56, 24, 32, 18, 18, 20, 16, 30, 30, 22]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        if rows:
            table = Table(displayName="MainTable", ref=f"A1:K{len(rows) + 1}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    @staticmethod
    def transfer_to_ote(ote_bytes: bytes, rows: list[dict[str, Any]], result_key: str) -> tuple[BytesIO, int]:
        if result_key not in RESULT_KEYS:
            raise ValueError("Choose a valid result column to transfer to OTE.")

        outcome_by_case = {
            int(row["testCaseId"]): str(row.get(result_key, "")).strip()
            for row in rows
            if row.get("testCaseId") and str(row.get(result_key, "")).strip()
        }
        if not outcome_by_case:
            raise ValueError("The selected result column has no outcomes to transfer.")

        # OTE workbooks can be large. Loading the full workbook with openpyxl can
        # expand memory usage dramatically, so update only worksheet XML inside the
        # XLSX ZIP package. This also preserves workbook formatting and structure.
        spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        relationships_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        ET.register_namespace("", spreadsheet_ns)
        ET.register_namespace("r", relationships_ns)

        def qname(name: str) -> str:
            return f"{{{spreadsheet_ns}}}{name}"

        def column_number(reference: str) -> int | None:
            match = re.match(r"([A-Z]+)", reference or "")
            if not match:
                return None
            value = 0
            for char in match.group(1):
                value = value * 26 + ord(char) - 64
            return value

        def column_letters(number: int) -> str:
            letters = ""
            while number:
                number, remainder = divmod(number - 1, 26)
                letters = chr(65 + remainder) + letters
            return letters

        def load_shared_strings(archive: ZipFile) -> list[str]:
            try:
                root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            except KeyError:
                return []
            values: list[str] = []
            for item in root.findall(qname("si")):
                values.append("".join(node.text or "" for node in item.iter(qname("t"))))
            return values

        def cell_text(cell: ET.Element | None, shared_strings: list[str]) -> str:
            if cell is None:
                return ""
            cell_type = cell.get("t")
            if cell_type == "inlineStr":
                return "".join(node.text or "" for node in cell.iter(qname("t")))
            value = cell.find(qname("v"))
            raw = value.text if value is not None and value.text is not None else ""
            if cell_type == "s":
                try:
                    return shared_strings[int(raw)]
                except (ValueError, IndexError):
                    return ""
            return raw

        def set_inline_text(cell: ET.Element, value: str) -> None:
            for child in list(cell):
                if child.tag in {qname("v"), qname("is"), qname("f")}:
                    cell.remove(child)
            cell.set("t", "inlineStr")
            inline = ET.SubElement(cell, qname("is"))
            text = ET.SubElement(inline, qname("t"))
            text.text = value

        def worksheet_paths(archive: ZipFile) -> list[str]:
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {
                rel.get("Id"): rel.get("Target")
                for rel in rels.findall(f"{{{package_rel_ns}}}Relationship")
            }
            paths: list[str] = []
            sheets = workbook.find(qname("sheets"))
            if sheets is None:
                return paths
            for sheet in sheets.findall(qname("sheet")):
                rel_id = sheet.get(f"{{{relationships_ns}}}id")
                target = targets.get(rel_id or "")
                if not target:
                    continue
                if target.startswith("/"):
                    paths.append(target.lstrip("/"))
                else:
                    paths.append(posixpath.normpath(posixpath.join("xl", target)))
            return paths

        def update_sheet(xml_bytes: bytes, shared_strings: list[str], updated_cases: set[int]) -> tuple[bytes, bool]:
            root = ET.fromstring(xml_bytes)
            sheet_data = root.find(qname("sheetData"))
            if sheet_data is None:
                return xml_bytes, False

            header_map: dict[str, int] = {}
            rows_xml = sheet_data.findall(qname("row"))
            header_row = rows_xml[0] if rows_xml else None
            if header_row is None:
                return xml_bytes, False

            for cell in header_row.findall(qname("c")):
                column = column_number(cell.get("r", ""))
                if column is not None:
                    header_map[cell_text(cell, shared_strings).strip()] = column

            case_col = header_map.get("TestCaseId")
            title_col = header_map.get("Title")
            outcome_col = header_map.get("Outcome")
            if not case_col or not outcome_col:
                return xml_bytes, False

            changed = False
            current_case_id: int | None = None
            current_outcome = ""

            for row in rows_xml[1:]:
                row_number = int(row.get("r") or 0)
                cells = {
                    column_number(cell.get("r", "")): cell
                    for cell in row.findall(qname("c"))
                }

                raw_case = cell_text(cells.get(case_col), shared_strings)
                raw_title = cell_text(cells.get(title_col), shared_strings) if title_col else ""
                try:
                    candidate_id = int(float(raw_case))
                except (TypeError, ValueError):
                    candidate_id = None

                is_case_header = candidate_id is not None and (
                    title_col is None or raw_title.strip() not in {"", "11"}
                )
                if is_case_header:
                    current_case_id = candidate_id
                    current_outcome = outcome_by_case.get(candidate_id, "")
                    if current_outcome:
                        updated_cases.add(candidate_id)

                if current_case_id is None or not current_outcome:
                    continue

                outcome_cell = cells.get(outcome_col)
                if outcome_cell is None:
                    outcome_cell = ET.Element(
                        qname("c"),
                        {"r": f"{column_letters(outcome_col)}{row_number}"},
                    )
                    inserted = False
                    for index, existing in enumerate(row.findall(qname("c"))):
                        existing_col = column_number(existing.get("r", "")) or 0
                        if existing_col > outcome_col:
                            row.insert(index, outcome_cell)
                            inserted = True
                            break
                    if not inserted:
                        row.append(outcome_cell)

                if cell_text(outcome_cell, shared_strings) != current_outcome:
                    set_inline_text(outcome_cell, current_outcome)
                    changed = True

            if not changed:
                return xml_bytes, False
            return ET.tostring(root, encoding="utf-8", xml_declaration=True), True

        try:
            source = BytesIO(ote_bytes)
            output = BytesIO()
            updated_cases: set[int] = set()
            with ZipFile(source, "r") as archive:
                shared_strings = load_shared_strings(archive)
                sheet_paths = set(worksheet_paths(archive))
                replacements: dict[str, bytes] = {}
                for path in sheet_paths:
                    try:
                        new_xml, changed = update_sheet(
                            archive.read(path), shared_strings, updated_cases
                        )
                    except KeyError:
                        continue
                    if changed:
                        replacements[path] = new_xml

                if not updated_cases:
                    raise ValueError(
                        "No matching TestCaseId values were found in the selected OTE workbook."
                    )

                with ZipFile(output, "w", compression=ZIP_DEFLATED) as rebuilt:
                    for info in archive.infolist():
                        rebuilt.writestr(info, replacements.get(info.filename, archive.read(info.filename)))
        except BadZipFile as error:
            raise ValueError("The selected OTE file is not a valid .xlsx workbook.") from error

        output.seek(0)
        return output, len(updated_cases)

