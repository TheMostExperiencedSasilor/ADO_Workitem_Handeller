import sys
from pathlib import Path
from unittest.mock import Mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from config import AppConfig
from services.ado_client import AdoClient
from services.test_assignment_workbook import (
    WORKBOOK_HEADERS,
    TestAssignmentWorkbookService,
)
from services.test_run_publisher import TestRunPublisher


def config(monkeypatch):
    monkeypatch.setenv("ADO_ORGANIZATION", "configured-org")
    monkeypatch.setenv("ADO_PROJECT", "configured project")
    monkeypatch.setenv("ADO_PAT", "unit-test-secret")
    return AppConfig.from_env()


def test_define_metadata_and_assigned_tester_join_keeps_point_ids(monkeypatch):
    client = AdoClient(config(monkeypatch))
    service = TestAssignmentWorkbookService(client)

    client._read_test_plan_pages = Mock(side_effect=[
        [
            {"workItem": {"id": 10}, "order": 2},
            {"workItem": {"id": 20}, "order": 1},
            {"workItem": {"id": 30}, "order": 3},
        ],
        [
            {
                "id": 501,
                "testCaseReference": {"id": 10},
                "configuration": {"id": "2", "name": "Windows 11"},
                "tester": {"displayName": "Jane QA", "uniqueName": "jane@example.test"},
                "results": {"outcome": "passed"},
            },
            {
                "id": 502,
                "testCaseReference": {"id": 10},
                "configuration": {"id": "3", "name": "Windows 10"},
                "tester": {"displayName": "Jane QA", "uniqueName": "jane@example.test"},
                "results": {"outcome": "passed"},
            },
            {
                "id": 600,
                "testCaseReference": {"id": 20},
                "tester": {"displayName": "Someone Else", "uniqueName": "other@example.test"},
                "results": {"outcome": "passed"},
            },
        ],
    ])
    client.read_work_items = Mock(return_value=[
        {
            "id": 10,
            "fields": {
                "System.Title": "Case Ten",
                "Custom.ProductArea": "Sample Files_HYSYS",
                "Custom.AutomationScriptName": "VSTS10.cs",
            },
        },
        {
            "id": 20,
            "fields": {
                "System.Title": "Case Twenty",
                "Custom.ProductArea": "Dynamics",
                "Custom.AutomationScriptName": "VSTS20.cs",
            },
        },
        {"id": 30, "fields": {"System.Title": "Case Thirty"}},
    ])

    result = service.build_assignment_rows(83602, 141923, "jane@example.test")

    assert result["defineCount"] == 3
    assert result["assignedCount"] == 1
    assert result["assignedPointCount"] == 2
    assert result["matchedTester"] == "Jane QA"
    row = result["rows"][0]
    assert row["testCaseId"] == 10
    assert row["title"] == "Case Ten"
    assert row["productArea"] == "Sample Files_HYSYS"
    assert row["automationScriptName"] == "VSTS10.cs"
    assert row["testPointIds"] == [501, 502]
    assert row["configurations"] == ["Windows 11", "Windows 10"]
    for key in (
        "round1Results", "round2Results", "singleRunResults", "manualRun",
        "comment", "solution", "defects",
    ):
        assert row[key] == ""


def test_ambiguous_partial_tester_is_rejected(monkeypatch):
    client = AdoClient(config(monkeypatch))
    service = TestAssignmentWorkbookService(client)
    client._read_test_plan_pages = Mock(side_effect=[
        [{"workItem": {"id": 1}, "order": 1}],
        [
            {"id": 1, "testCaseReference": {"id": 1}, "tester": {"displayName": "Jane One"}},
            {"id": 2, "testCaseReference": {"id": 1}, "tester": {"displayName": "Jane Two"}},
        ],
    ])
    client.read_work_items = Mock(return_value=[{"id": 1, "fields": {"System.Title": "Case"}}])

    try:
        service.build_assignment_rows(1, 2, "Jane")
        assert False, "Expected ambiguous tester search to fail"
    except ValueError as error:
        assert "multiple people" in str(error)


def test_excel_exports_current_tracking_values():
    rows = [{
        "testCaseId": 10,
        "title": "Case Ten",
        "productArea": "Area",
        "automationScriptName": "VSTS10.cs",
        "round1Results": "Passed",
        "round2Results": "Failed",
        "singleRunResults": "",
        "manualRun": "N/A",
        "comment": "note",
        "solution": "fix",
        "defects": "12345",
    }]

    stream = TestAssignmentWorkbookService.make_workbook(rows)

    from openpyxl import load_workbook
    workbook = load_workbook(stream)
    assert workbook.sheetnames == ["Main"]
    sheet = workbook["Main"]
    assert [cell.value for cell in sheet[1]] == WORKBOOK_HEADERS
    assert [sheet[f"{column}2"].value for column in "ABCDEFGHIJK"] == [
        10, "Case Ten", "Area", "VSTS10.cs", "Passed", "Failed", None,
        "N/A", "note", "fix", "12345",
    ]
    assert sheet.freeze_panes == "A2"


def test_publish_creates_run_updates_point_results_and_completes_it():
    client = Mock()
    client.read_test_points.return_value = [
        {"testPointId": 501, "testCaseId": 10},
        {"testPointId": 502, "testCaseId": 10},
    ]
    client.create_test_run.return_value = {
        "id": 77,
        "name": "Round 1 Jane",
        "webAccessUrl": "https://dev.azure.com/org/project/_testManagement/runs/77",
    }
    client.list_test_results.return_value = [
        {"id": 9001, "testPoint": {"id": "501"}},
        {"id": 9002, "testPoint": {"id": "502"}},
    ]
    client.update_test_results.return_value = []
    client.update_test_run.return_value = {
        "id": 77,
        "name": "Round 1 Jane",
        "state": "Completed",
        "webAccessUrl": "https://dev.azure.com/org/project/_testManagement/runs/77",
    }

    result = TestRunPublisher(client).publish(
        plan_id=83602,
        suite_id=141923,
        result_key="round1Results",
        run_name="Round 1 Jane",
        rows=[{
            "testCaseId": 10,
            "testPointIds": [501, 502],
            "round1Results": "Passed",
            "comment": "good",
            "solution": "",
            "defects": "",
        }],
    )

    client.create_test_run.assert_called_once_with(83602, "Round 1 Jane", [501, 502], automated=True)
    client.update_test_results.assert_called_once_with(77, [
        {"id": 9001, "outcome": "Passed", "state": "Completed", "comment": "good"},
        {"id": 9002, "outcome": "Passed", "state": "Completed", "comment": "good"},
    ])
    client.update_test_run.assert_called_once_with(77, {"state": "Completed"})
    assert result["runId"] == 77
    assert result["publishedPoints"] == 2
    assert result["publishedCases"] == 1
    assert result["outcomes"] == {"Passed": 2}


def test_publish_manual_run_maps_outcome_and_marks_run_manual():
    client = Mock()
    client.read_test_points.return_value = [{"testPointId": 5, "testCaseId": 1}]
    client.create_test_run.return_value = {"id": 8, "name": "Manual"}
    client.list_test_results.return_value = [{"id": 80, "testPoint": {"id": "5"}}]
    client.update_test_results.return_value = []
    client.update_test_run.return_value = {"id": 8, "name": "Manual", "state": "Completed"}

    TestRunPublisher(client).publish(
        1, 2,
        [{"testCaseId": 1, "testPointIds": [5], "manualRun": "N/A"}],
        "manualRun",
        "Manual",
    )

    client.create_test_run.assert_called_once_with(1, "Manual", [5], automated=False)
    client.update_test_results.assert_called_once_with(8, [
        {"id": 80, "outcome": "NotApplicable", "state": "Completed"},
    ])


def test_publish_rejects_old_session_without_test_point_ids():
    client = Mock()
    publisher = TestRunPublisher(client)
    try:
        publisher.publish(
            1, 2,
            [{"testCaseId": 1, "round1Results": "Passed"}],
            "round1Results",
            "Old Session",
        )
        assert False, "Expected missing point mapping to fail"
    except ValueError as error:
        assert "Reload the assignment" in str(error)
    client.create_test_run.assert_not_called()
