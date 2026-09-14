import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from openpyxl import Workbook, load_workbook

from services.test_assignment_workbook import TestAssignmentWorkbookService


def make_ote_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "83602;V15.2_HYSYS"
    sheet.append([
        "TestCaseId", "Title", "TestStep", "StepAction", "StepExpected",
        "TestPointId", "Configuration", "Tester", "Outcome", "Comment",
    ])
    sheet.append([24153, "Case A", "11", "11", "11", "200291:0", "Windows 10", "Jane QA", "", ""])
    sheet.append(["11", "11", "1", "Do A", "Expected A", "11", "11", "11", "", "11"])
    sheet.append(["11", "11", "2", "Do B", "Expected B", "11", "11", "11", "", "11"])
    sheet.append([24157, "Case B", "11", "11", "11", "200293:0", "Windows 10", "Jane QA", "", ""])
    sheet.append(["11", "11", "1", "Do C", "Expected C", "11", "11", "11", "", "11"])
    raw = BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def test_transfer_to_ote_fills_matching_blocks_without_changing_tester():
    rows = [
        {"testCaseId": 24153, "round1Results": "Passed"},
        {"testCaseId": 24157, "round1Results": "Failed"},
    ]

    output, updated = TestAssignmentWorkbookService.transfer_to_ote(
        make_ote_workbook(), rows, "round1Results"
    )
    sheet = load_workbook(output)["83602;V15.2_HYSYS"]

    assert updated == 2
    assert [sheet[f"I{row}"].value for row in range(2, 7)] == [
        "Passed", "Passed", "Passed", "Failed", "Failed",
    ]
    assert sheet["H2"].value == "Jane QA"
    assert sheet["H5"].value == "Jane QA"


def test_transfer_to_ote_requires_logged_results():
    try:
        TestAssignmentWorkbookService.transfer_to_ote(
            make_ote_workbook(), [{"testCaseId": 24153, "round1Results": ""}], "round1Results"
        )
        assert False, "Expected blank OTE result source to be rejected"
    except ValueError as error:
        assert "no outcomes" in str(error).lower()
